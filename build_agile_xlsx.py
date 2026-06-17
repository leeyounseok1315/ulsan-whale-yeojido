#!/usr/bin/env python3
"""울산고래여지도 애자일 산출물 엑셀 생성기.
워크플로가 만든 data.json을 읽어 새 .xlsx(여지도 세피아 톤)로 출력한다.
원본 템플릿(AgileProcess_XXX팀 (1).xlsx)은 건드리지 않는다. BurndownCharts 시트는 제외.
"""
import json
import sys
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = "/Users/name1315/Desktop/2026관광데이터"
DATA = sys.argv[1] if len(sys.argv) > 1 else f"{ROOT}/_agile_data.json"
OUT = f"{ROOT}/울산고래여지도_AgileProcess.xlsx"

with open(DATA, encoding="utf-8") as f:
    d = json.load(f)

# ---- 여지도 세피아 팔레트 ----
SEPIA_DARK = "5C4326"     # 제목 띠
SEPIA_HEAD = "8C6D46"     # 헤더
SEPIA_SUB = "C8A97E"      # 서브 헤더
SEPIA_LITE = "EFE6D5"     # 옅은 배경
WHITE = "FFFFFF"

thin = Side(style="thin", color="B49A6B")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(bold=False, color="2B2117", size=11):
    return Font(name="맑은 고딕", bold=bold, color=color, size=size)

def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

A_WRAP_TOP = Alignment(wrap_text=True, vertical="top", horizontal="left")
A_CTR = Alignment(wrap_text=True, vertical="center", horizontal="center")
A_CTR_TOP = Alignment(wrap_text=True, vertical="top", horizontal="center")

def style_cell(ws, coord, value, *, bold=False, fg=None, color="2B2117",
               align=A_WRAP_TOP, border=True, size=11):
    c = ws[coord]
    c.value = value
    c.font = font(bold=bold, color=color, size=size)
    if fg:
        c.fill = fill(fg)
    c.alignment = align
    if border:
        c.border = BORDER
    return c

wb = Workbook()

# =========================================================
# 1) System Overview
# =========================================================
ws = wb.active
ws.title = "System Overview"
ws.sheet_view.showGridLines = False
so = d["systemOverview"]

ws.column_dimensions["A"].width = 3
for col, w in zip("BCDEF", [16, 26, 26, 26, 22]):
    ws.column_dimensions[col].width = w

style_cell(ws, "B2", "울산고래여지도 — 애자일 개발 산출물", bold=True,
           fg=SEPIA_DARK, color=WHITE, align=A_CTR, size=14)
ws.merge_cells("B2:F2")
ws.row_dimensions[2].height = 30

style_cell(ws, "B4", "1. 시스템 개요", bold=True, fg=SEPIA_HEAD, color=WHITE)
ws.merge_cells("B4:F4")
style_cell(ws, "B5", so["overview"], align=A_WRAP_TOP)
ws.merge_cells("B5:F5")
ws.row_dimensions[5].height = 92

style_cell(ws, "B7", "2. 시스템 목적", bold=True, fg=SEPIA_HEAD, color=WHITE)
ws.merge_cells("B7:F7")
style_cell(ws, "B8", so["purpose"], align=A_WRAP_TOP)
ws.merge_cells("B8:F8")
ws.row_dimensions[8].height = 88

style_cell(ws, "B10", "3. 팀 구성원 및 역할", bold=True, fg=SEPIA_HEAD, color=WHITE)
ws.merge_cells("B10:F10")
style_cell(ws, "B11", "학번", bold=True, fg=SEPIA_SUB, align=A_CTR)
style_cell(ws, "C11", "이름", bold=True, fg=SEPIA_SUB, align=A_CTR)
style_cell(ws, "D11", "역할", bold=True, fg=SEPIA_SUB, align=A_CTR)
ws.merge_cells("D11:F11")
r = 12
for m in so["members"]:
    style_cell(ws, f"B{r}", m["studentId"], align=A_CTR)
    style_cell(ws, f"C{r}", m["name"], align=A_CTR)
    style_cell(ws, f"D{r}", m["role"], align=A_WRAP_TOP)
    ws.merge_cells(f"D{r}:F{r}")
    ws.row_dimensions[r].height = 22
    r += 1

# =========================================================
# 2) Process Definition  (스프린트 = 컬럼)
# =========================================================
ws = wb.create_sheet("Process Definition")
ws.sheet_view.showGridLines = False
sprints = sorted(d["processDefinition"]["sprints"], key=lambda s: s["sprintNo"])

style_cell(ws, "B2", "Process Definition (스프린트 정의)", bold=True,
           fg=SEPIA_DARK, color=WHITE, align=A_CTR, size=13)
last_col = get_column_letter(2 + len(sprints))  # B + n sprints
ws.merge_cells(f"B2:{last_col}2")
ws.row_dimensions[2].height = 26

rows = [
    ("Date (시작일)", "startDate"),
    ("Sprint No.", "sprintNo"),
    ("기간(주차)", "weeks"),
    ("Backlog No.", "backlogNos"),
    ("Minor Milestone", "minorMilestone"),
    ("Major Milestone", "majorMilestone"),
]
ws.column_dimensions["B"].width = 18
for i, label_key in enumerate(rows):
    label = label_key[0]
    rr = 3 + i
    style_cell(ws, f"B{rr}", label, bold=True, fg=SEPIA_HEAD, color=WHITE,
               align=Alignment(vertical="center", horizontal="left", wrap_text=True))

for j, sp in enumerate(sprints):
    col = get_column_letter(3 + j)
    ws.column_dimensions[col].width = 22
    for i, (label, key) in enumerate(rows):
        rr = 3 + i
        val = sp[key]
        if key == "sprintNo":
            val = f"Sprint {val}"
        is_head = key in ("startDate", "sprintNo")
        style_cell(ws, f"{col}{rr}", val,
                   bold=is_head,
                   fg=(SEPIA_SUB if is_head else (SEPIA_LITE if key == "majorMilestone" and val else None)),
                   align=A_CTR_TOP if key in ("startDate", "sprintNo", "weeks", "backlogNos") else A_WRAP_TOP)
ws.row_dimensions[3 + 4].height = 66  # Minor
ws.row_dimensions[3 + 5].height = 52  # Major
ws.freeze_panes = "C3"

# =========================================================
# 3) Product Backlog
# =========================================================
ws = wb.create_sheet("Product Backlog")
ws.sheet_view.showGridLines = False
items = sorted(d["productBacklog"]["items"], key=lambda x: x["backlogNo"])

heads = ["Backlog\nNO.", "메인 메뉴", "Sub-1 메뉴", "우선\n순위", "Demo Scenario (데모 시작~종료 사용자 행위 그대로)"]
widths = [9, 20, 22, 9, 80]
for i, (h, w) in enumerate(zip(heads, widths)):
    col = get_column_letter(1 + i)
    ws.column_dimensions[col].width = w
    style_cell(ws, f"{col}2", h, bold=True, fg=SEPIA_HEAD, color=WHITE, align=A_CTR)
ws.row_dimensions[2].height = 32

PRIO_FILL = {"MVP": "C8A97E", "Phase2": "DCD0B8", "Stretch": "EFE6D5"}
r = 3
for it in items:
    style_cell(ws, f"A{r}", it["backlogNo"], align=A_CTR_TOP, bold=True)
    style_cell(ws, f"B{r}", it["mainMenu"], align=A_WRAP_TOP)
    style_cell(ws, f"C{r}", it["subMenu"], align=A_WRAP_TOP)
    style_cell(ws, f"D{r}", it["priority"], align=A_CTR_TOP,
               fg=PRIO_FILL.get(it["priority"]), bold=True)
    style_cell(ws, f"E{r}", it["demoScenario"], align=A_WRAP_TOP)
    n_lines = it["demoScenario"].count("\n") + 1
    ws.row_dimensions[r].height = max(40, min(220, 16 * n_lines + 8))
    r += 1
ws.freeze_panes = "A3"

# =========================================================
# 4) Sprint Backlog  (스프린트별 행 + Sprint No. 병합)
# =========================================================
ws = wb.create_sheet("Sprint Backlog")
ws.sheet_view.showGridLines = False
sbrows = sorted(d["sprintBacklog"]["rows"], key=lambda x: (x["sprintNo"], x["backlogNo"]))

heads = ["Sprint\nNo.", "BackLog", "개발 리더·발표자\n(이윤석 / 백엔드)",
         "개발자 2\n(이숙빈 / 프론트엔드)", "개발자 3\n(해당 없음)", "결과"]
widths = [8, 26, 40, 40, 13, 10]
for i, (h, w) in enumerate(zip(heads, widths)):
    col = get_column_letter(2 + i)  # B..G
    ws.column_dimensions[col].width = w
    style_cell(ws, f"{col}2", h, bold=True, fg=SEPIA_HEAD, color=WHITE, align=A_CTR)
ws.row_dimensions[2].height = 36

RESULT_FILL = {"완료": "C8A97E", "진행중": "DCD0B8", "예정": "EFE6D5"}
r = 3
sprint_spans = {}  # sprintNo -> [start_row, end_row]
for row in sbrows:
    sn = row["sprintNo"]
    style_cell(ws, f"B{r}", f"Sprint {sn}", align=A_CTR, bold=True, fg=SEPIA_SUB)
    style_cell(ws, f"C{r}", f'{row["backlogNo"]}. {row["backlogName"]}', align=A_WRAP_TOP, bold=True)
    style_cell(ws, f"D{r}", row["leaderTask"], align=A_WRAP_TOP)
    style_cell(ws, f"E{r}", row["dev2Task"], align=A_WRAP_TOP)
    style_cell(ws, f"F{r}", "-", align=A_CTR_TOP)
    style_cell(ws, f"G{r}", row["result"], align=A_CTR_TOP,
               fg=RESULT_FILL.get(row["result"]), bold=True)
    longest = max(row["leaderTask"].count("\n"), row["dev2Task"].count("\n")) + 1
    txt_len = max(len(row["leaderTask"]), len(row["dev2Task"]))
    ws.row_dimensions[r].height = max(40, min(180, 16 * longest + (txt_len // 38) * 16 + 8))
    sprint_spans.setdefault(sn, [r, r])[1] = r
    r += 1
# Sprint No. 셀 병합
for sn, (a, b) in sprint_spans.items():
    if b > a:
        ws.merge_cells(f"B{a}:B{b}")
ws.freeze_panes = "B3"

wb.save(OUT)

print("SAVED:", OUT)
print("sheets:", wb.sheetnames)
print("backlog items:", len(items))
print("sprints:", len(sprints))
print("sprint-backlog rows:", len(sbrows))
if d.get("issuesFound"):
    print("issuesFound:")
    for s in d["issuesFound"]:
        print("  -", s)
